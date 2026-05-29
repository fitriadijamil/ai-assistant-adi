import json, re, sys
sys.path.insert(0, '/tmp/xlvenv/lib/python3.14/site-packages')
import openpyxl

wb = openpyxl.load_workbook('/Users/macbookair/Desktop/MEI_2026_MD_HIVIEW_PRICELIST.xlsx', read_only=True, data_only=True)

def safe_int(v):
    if v is None: return None
    try: return int(float(str(v).replace(',','').strip()))
    except: return None

def clean(v):
    if v is None: return ''
    s = str(v).strip()
    return s

def get_price(row, idx):
    v = safe_int(row[idx]) if idx < len(row) else None
    return v

# ============================================================
# PARSE CAMERAS
# ============================================================
cameras = []

# --- IPC D-SERIES (IP cameras) ---
# cols: TIPE(0), LENSA(1), GAMBAR(2), RESOLUSI(3), KETERANGAN(4), MSRP(5), MD PRICE(6), NON MD(7), ONLINE(8)
ws = wb['IPC D-SERIES']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    tipe = clean(row[0]) if row[0] else ''
    if not tipe: continue
    resolusi_raw = clean(row[3]) if len(row) > 3 else ''
    resolusi_map = {'2MP':'2MP','1080P':'2MP','3K':'5MP','5MP':'5MP','4MP':'4MP','8MP':'8MP'}
    resolusi = '2MP'
    for k,v in resolusi_map.items():
        if k in resolusi_raw.upper() or k in resolusi_raw: resolusi = v
    harga = get_price(row, 6)  # MD PRICE
    if not harga: harga = get_price(row, 5)  # MSRP fallback
    cam = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "tipe": "IP",
        "jenis": "Turret" if "T-" in tipe or "T " in tipe or tipe.startswith("TH-T") else "Bullet",
        "resolusi": resolusi,
        "poe": True,
        "bitrate_default_bps": {"2MP":4000000,"3MP":6000000,"4MP":8000000,"5MP":10000000,"8MP":16000000}.get(resolusi, 8000000),
        "fitur": ["IP camera"],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    cameras.append(cam)

# --- IPC H-SERIES (need to read) ---
ws = wb['IPC H-SERIES']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    tipe = clean(row[0]) if row[0] else ''
    if not tipe: continue
    resolusi_raw = clean(row[3]) if len(row) > 3 else ''
    resolusi = '2MP'
    resolusi_map = {'2MP':'2MP','1080P':'2MP','3K':'5MP','5MP':'5MP','4MP':'4MP','8MP':'8MP'}
    for k,v in resolusi_map.items():
        if k in resolusi_raw.upper() or k in resolusi_raw: resolusi = v
    harga = get_price(row, 6)
    if not harga: harga = get_price(row, 5)
    lens = clean(row[1]) if len(row) > 1 else ''
    cam = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "tipe": "IP",
        "jenis": "Turret" if "T-" in tipe or "T " in tipe or tipe.startswith("TH-T") else "Bullet",
        "resolusi": resolusi,
        "poe": True,
        "bitrate_default_bps": {"2MP":4000000,"3MP":6000000,"4MP":8000000,"5MP":10000000,"8MP":16000000}.get(resolusi, 8000000),
        "fitur": ["IP camera"],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    cameras.append(cam)

# --- C-SERIES CAM (Analog) ---
ws = wb['C-SERIES CAM']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:12]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    resolusi_raw = clean(row[4]) if len(row) > 4 else ''
    resolusi = '2MP'
    if '1080P' in resolusi_raw.upper(): resolusi = '2MP'
    lens = clean(row[2]) if len(row) > 2 else ''
    harga = get_price(row, 7)  # MD PRICE
    if not harga: harga = get_price(row, 6)
    jenis = "Bullet" if "B" in tipe else "Turret" if "T" in tipe else "Camera"
    cam = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "tipe": "Analog",
        "jenis": jenis,
        "resolusi": resolusi,
        "poe": False,
        "bitrate_default_bps": 0,
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 8),
        "harga_online": get_price(row, 9),
        "harga_msrp": get_price(row, 6)
    }
    cameras.append(cam)

# --- D-SERIES CAM (Analog) ---
ws = wb['D-SERIES CAM']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:12]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    resolusi_raw = clean(row[4]) if len(row) > 4 else ''
    resolusi = '2MP'
    if '1080P' in resolusi_raw.upper(): resolusi = '2MP'
    harga = get_price(row, 7)
    if not harga: harga = get_price(row, 6)
    jenis = "Bullet" if "B" in tipe else "Turret" if "T" in tipe else "Camera"
    if "DOME" in seri.upper(): jenis = "Dome"
    cam = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "tipe": "Analog",
        "jenis": jenis,
        "resolusi": resolusi,
        "poe": False,
        "bitrate_default_bps": 0,
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 8),
        "harga_online": get_price(row, 9),
        "harga_msrp": get_price(row, 6)
    }
    cameras.append(cam)

# --- H-SERIES CAM (Analog) ---
ws = wb['H-SERIES CAM']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:12]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    resolusi_raw = clean(row[4]) if len(row) > 4 else ''
    resolusi = '2MP'
    if '1080P' in resolusi_raw.upper(): resolusi = '2MP'
    if '3K' in resolusi_raw or '5MP' in resolusi_raw: resolusi = '5MP'
    harga = get_price(row, 7)
    if not harga: harga = get_price(row, 6)
    jenis = "Bullet" if "B" in tipe else "Turret" if "T" in tipe else "Camera"
    cam = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "tipe": "Analog",
        "jenis": jenis,
        "resolusi": resolusi,
        "poe": False,
        "bitrate_default_bps": 0,
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 8),
        "harga_online": get_price(row, 9),
        "harga_msrp": get_price(row, 6)
    }
    cameras.append(cam)

with open('/Users/macbookair/Desktop/ai-assistant-adi/backend/app/catalog/hiview_cameras.json', 'w') as f:
    json.dump(cameras, f, indent=2, ensure_ascii=False)
print(f"Cameras: {len(cameras)} products")

# ============================================================
# PARSE NVR / XVR
# ============================================================
nvrs = []

# --- D- SERIES XVR ---
ws = wb['D- SERIES XVR']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    ch_raw = clean(row[2]) if len(row) > 2 else ''
    ch = 0
    m = re.search(r'(\d+)', ch_raw)
    if m: ch = int(m.group(1))
    harga = get_price(row, 6)  # MD
    if not harga: harga = get_price(row, 5)
    rec = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "channel": ch,
        "max_resolusi": "5MP",
        "max_hdd_tb": 6,
        "poe_builtin": False,
        "tipe": "XVR",
        "fitur": ["Hybrid"],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    nvrs.append(rec)

# --- H-SERIES XVR ---
ws = wb['H-SERIES XVR']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    ch_raw = clean(row[2]) if len(row) > 2 else ''
    ch = 0
    m = re.search(r'(\d+)', ch_raw)
    if m: ch = int(m.group(1))
    harga = get_price(row, 6)
    if not harga: harga = get_price(row, 5)
    rec_type = "XVR"
    if "DVR" in seri.upper(): rec_type = "DVR"
    rec = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "channel": ch,
        "max_resolusi": "5MP",
        "max_hdd_tb": 6,
        "poe_builtin": False,
        "tipe": rec_type,
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    nvrs.append(rec)

# --- NVR D-SERIES ---
ws = wb['NVR D-SERIES']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    ch_raw = clean(row[2]) if len(row) > 2 else ''
    ch = 0
    m = re.search(r'(\d+)', ch_raw)
    if m: ch = int(m.group(1))
    harga = get_price(row, 6)
    if not harga: harga = get_price(row, 5)
    rec = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "channel": ch,
        "max_resolusi": "8MP",
        "max_hdd_tb": 10,
        "poe_builtin": "POE" in seri.upper() or "P" in tipe.upper(),
        "tipe": "NVR",
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    nvrs.append(rec)

# --- NVR H- SERIES ---
ws = wb['NVR H- SERIES']
headers = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))[:10]]
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if row[1] else ''
    if not tipe: continue
    ch_raw = clean(row[2]) if len(row) > 2 else ''
    ch = 0
    m = re.search(r'(\d+)', ch_raw)
    if m: ch = int(m.group(1))
    harga = get_price(row, 6)
    if not harga: harga = get_price(row, 5)
    rec = {
        "id": "hv-" + re.sub(r'[^a-zA-Z0-9]', '', tipe.lower()),
        "nama": tipe,
        "brand": "Hiview",
        "channel": ch,
        "max_resolusi": "8MP",
        "max_hdd_tb": 10,
        "poe_builtin": "POE" in seri.upper() or "P" in tipe.upper(),
        "tipe": "NVR",
        "fitur": [],
        "harga_md": harga,
        "harga_non_md": get_price(row, 7),
        "harga_online": get_price(row, 8),
        "harga_msrp": get_price(row, 5)
    }
    nvrs.append(rec)

with open('/Users/macbookair/Desktop/ai-assistant-adi/backend/app/catalog/hiview_nvr.json', 'w') as f:
    json.dump(nvrs, f, indent=2, ensure_ascii=False)
print(f"NVR/XVR: {len(nvrs)} products")

# ============================================================
# PARSE ACCESSORIES
# ============================================================
accessories = []

ws = wb['AKSESORIS']
for row in ws.iter_rows(min_row=2, values_only=True):
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    seri = clean(row[0]) if row[0] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 5) if len(row) > 5 else None  # MD PRICE
    if not harga: harga = get_price(row, 4) if len(row) > 4 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Accessories",
        "harga_md": harga,
        "harga_non_md": get_price(row, 6) if len(row) > 6 else None,
        "harga_online": get_price(row, 7) if len(row) > 7 else None,
        "harga_msrp": get_price(row, 4) if len(row) > 4 else None
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

# WIRELESS
ws = wb['WIRELESS']
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 6) if len(row) > 6 else None
    if not harga: harga = get_price(row, 5) if len(row) > 5 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Wireless",
        "harga_md": harga
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

# TRANSMISSION (switches, etc)
ws = wb['TRANSMISSION']
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 6) if len(row) > 6 else None
    if not harga: harga = get_price(row, 5) if len(row) > 5 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Transmission",
        "harga_md": harga
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

# MONITOR
ws = wb['MONITOR']
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 6) if len(row) > 6 else None
    if not harga: harga = get_price(row, 5) if len(row) > 5 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Monitor",
        "harga_md": harga
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

# KABEL
ws = wb['KABEL']
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 5) if len(row) > 5 else None
    if not harga: harga = get_price(row, 4) if len(row) > 4 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Cable",
        "harga_md": harga
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

# ACCESS CONTROL
ws = wb['ACCESS CONTROL']
for row in ws.iter_rows(min_row=2, values_only=True):
    seri = clean(row[0]) if row[0] else ''
    tipe = clean(row[1]) if len(row) > 1 and row[1] else ''
    if not tipe and not seri: continue
    if not tipe: tipe = seri
    harga = get_price(row, 6) if len(row) > 6 else None
    if not harga: harga = get_price(row, 5) if len(row) > 5 else None
    acc = {
        "nama": tipe,
        "brand": "Hiview",
        "kategori": "Access Control",
        "harga_md": harga
    }
    if seri and seri != tipe: acc["seri"] = seri
    accessories.append(acc)

with open('/Users/macbookair/Desktop/ai-assistant-adi/backend/app/catalog/hiview_accessories.json', 'w') as f:
    json.dump(accessories, f, indent=2, ensure_ascii=False)
print(f"Accessories: {len(accessories)} products")
print("Done!")
